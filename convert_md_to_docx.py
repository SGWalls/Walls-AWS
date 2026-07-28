#!/usr/bin/env python3
"""Convert markdown file to Word document with formatting."""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_LINE_SPACING


def convert_markdown_to_docx(md_file, docx_file):
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Skip horizontal rules
        if line.strip() == '---':
            i += 1
            continue
        
        # H1 - Main title
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # H2 - Sections
        elif line.startswith('## '):
            doc.add_paragraph()  # Add space before section
            p = doc.add_heading(line[3:], level=2)
        
        # H3 - Subsections
        elif line.startswith('### '):
            p = doc.add_heading(line[4:], level=3)
        
        # Bullet points
        elif line.startswith('- '):
            text = line[2:]
            # Check if bold
            if text.startswith('**') and ':**' in text:
                p = doc.add_paragraph(style='List Bullet')
                bold_end = text.index(':**') + 2
                run = p.add_run(text[2:bold_end])
                run.bold = True
                p.add_run(text[bold_end:])
            else:
                doc.add_paragraph(text, style='List Bullet')
        
        # Nested bullet points
        elif line.startswith('  - '):
            text = line[4:]
            p = doc.add_paragraph(text, style='List Bullet 2')
        
        # Empty line
        elif line.strip() == '':
            pass  # Skip empty lines, spacing handled by paragraphs
        
        # Regular paragraph
        else:
            if line.strip():
                p = doc.add_paragraph(line)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        
        if '**' in line:
            p = doc.add_paragraph()
            parts = re.split(r'\*\*(.*?)\*\*', line)
            for i, part in enumerate(parts):
                if i % 2 == 0:
                    p.add_run(part)
                else:
                    p.add_run(part).bold = True

        i += 1
    
    doc.save(docx_file)
    print(f"Word document created: {docx_file}")

if __name__ == '__main__':
    md_file = r'c:\Users\sgwalls\Documents\AWS_Projects\Scripts\Python\Git_repository\Walls-AWS\Identity_Center_Migration_Glossary.md'
    docx_file = r'c:\Users\sgwalls\Documents\AWS_Projects\Scripts\Python\Git_repository\Walls-AWS\Identity_Center_Migration_Glossary.docx'
    
    convert_markdown_to_docx(md_file, docx_file)
