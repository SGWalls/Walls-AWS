import boto3
import logging
import os
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'helpers'))
from retriever import get_org_account_list, validate_sso_token, create_logger

logger = logging.getLogger(__name__)

target_accounts = [
    {
        "Id": "487692780423",
        "Name": "PRD_AO_PlanetALTiG"
    },
    {
        "Id": "424988161273",
        "Name": "STG_AO_PlanetALTiG"
    },
    {
        "Id": "294253399540",
        "Name": "SS_AO_PlanetALTiG"
    },
    {
        "Id": "461906805688",
        "Name": "POC_AO_PlanetALTiG"
    }
]
ROLE_NAME = 'AWSControlTowerExecution'
SESSION_NAME = 'ResourceDiscovery'


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def get_regions(session_or_client):
    """Return all opted-in EC2 regions."""
    ec2 = session_or_client if hasattr(session_or_client, 'describe_regions') \
        else session_or_client.client('ec2', region_name='us-east-1')
    return [r['RegionName'] for r in ec2.describe_regions(
        Filters=[{'Name': 'opt-in-status', 'Values': ['opt-in-not-required', 'opted-in']}]
    )['Regions']]


def _tag_name(tags):
    if not tags:
        return ''
    return next((t['Value'] for t in tags if t['Key'] == 'Name'), '')


def _make_client(credentials, service, region):
    return boto3.client(
        service,
        aws_access_key_id=credentials['AccessKeyId'],
        aws_secret_access_key=credentials['SecretAccessKey'],
        aws_session_token=credentials['SessionToken'],
        region_name=region,
    )


# ---------------------------------------------------------------------------
# Resource discovery functions
# ---------------------------------------------------------------------------

def get_vpcs(credentials, regions):
    rows = []
    for region in regions:
        try:
            ec2 = _make_client(credentials, 'ec2', region)
            for page in ec2.get_paginator('describe_vpcs').paginate():
                for vpc in page['Vpcs']:
                    rows.append({
                        'Region': region,
                        'Name': _tag_name(vpc.get('Tags')),
                        'VPC ID': vpc['VpcId'],
                    })
        except Exception as e:
            logger.debug(f"VPC {region}: {e}")
    return rows


def get_ec2_instances(credentials, regions):
    rows = []
    for region in regions:
        try:
            ec2 = _make_client(credentials, 'ec2', region)
            for page in ec2.get_paginator('describe_instances').paginate():
                for res in page['Reservations']:
                    for inst in res['Instances']:
                        rows.append({
                            'Region': region,
                            'Name': _tag_name(inst.get('Tags')),
                            'Instance Type': inst.get('InstanceType', ''),
                            'Availability Zone': inst.get('Placement', {}).get('AvailabilityZone', ''),
                            'Platform Details': inst.get('PlatformDetails', ''),
                        })
        except Exception as e:
            logger.debug(f"EC2 {region}: {e}")
    return rows


def get_load_balancers(credentials, regions):
    rows = []
    for region in regions:
        try:
            elb = _make_client(credentials, 'elbv2', region)
            for page in elb.get_paginator('describe_load_balancers').paginate():
                for lb in page['LoadBalancers']:
                    rows.append({
                        'Region': region,
                        'Name': lb.get('LoadBalancerName', ''),
                        'State': lb.get('State', {}).get('Code', ''),
                        'Type': lb.get('Type', ''),
                        'VPC ID': lb.get('VpcId', ''),
                        'Availability Zones': ', '.join(
                            az['ZoneName'] for az in lb.get('AvailabilityZones', [])
                        ),
                    })
        except Exception as e:
            logger.debug(f"ELB {region}: {e}")
    return rows


def get_s3_buckets(credentials):
    rows = []
    try:
        s3 = boto3.client(
            's3',
            aws_access_key_id=credentials['AccessKeyId'],
            aws_secret_access_key=credentials['SecretAccessKey'],
            aws_session_token=credentials['SessionToken'],
        )
        buckets = s3.list_buckets().get('Buckets', [])
        for bucket in buckets:
            try:
                region = s3.get_bucket_location(Bucket=bucket['Name'])['LocationConstraint'] or 'us-east-1'
            except Exception:
                region = 'unknown'
            rows.append({'Name': bucket['Name'], 'AWS Region': region})
    except Exception as e:
        logger.debug(f"S3: {e}")
    return rows


def get_elasticache(credentials, regions):
    rows = []
    for region in regions:
        try:
            ec = _make_client(credentials, 'elasticache', region)
            for page in ec.get_paginator('describe_cache_clusters').paginate(ShowCacheNodeInfo=False):
                for cluster in page['CacheClusters']:
                    rows.append({
                        'Region': region,
                        'Cache Name': cluster.get('CacheClusterId', ''),
                        'Status': cluster.get('CacheClusterStatus', ''),
                        'Engine Version': f"{cluster.get('Engine','')} {cluster.get('EngineVersion','')}",
                        'Configuration': cluster.get('CacheNodeType', ''),
                    })
        except Exception as e:
            logger.debug(f"ElastiCache {region}: {e}")
    return rows


def get_rds(credentials, regions):
    rows = []
    for region in regions:
        try:
            rds = _make_client(credentials, 'rds', region)
            for page in rds.get_paginator('describe_db_instances').paginate():
                for db in page['DBInstances']:
                    rows.append({
                        'Region': region,
                        'DB Identifier': db.get('DBInstanceIdentifier', ''),
                        'Status': db.get('DBInstanceStatus', ''),
                        'Role': 'Replica' if db.get('ReadReplicaSourceDBInstanceIdentifier') else 'Primary',
                        'Engine': f"{db.get('Engine','')} {db.get('EngineVersion','')}",
                        'Region & AZ': db.get('AvailabilityZone', ''),
                        'Size': db.get('DBInstanceClass', ''),
                        'Multi-AZ': db.get('MultiAZ', False),
                    })
        except Exception as e:
            logger.debug(f"RDS {region}: {e}")
    return rows


def get_acm_certificates(credentials, regions):
    rows = []
    for region in regions:
        try:
            acm = _make_client(credentials, 'acm', region)
            for page in acm.get_paginator('list_certificates').paginate():
                for cert in page['CertificateSummaryList']:
                    cert_id = cert['CertificateArn'].split('/')[-1]
                    rows.append({
                        'Region': region,
                        'Certificate ID': cert_id,
                        'Domain Name': cert.get('DomainName', ''),
                        'Type': cert.get('Type', ''),
                        'Status': cert.get('Status', ''),
                        'In Use': cert.get('InUse', ''),
                        'Renewal Eligibility': cert.get('RenewalEligibility', ''),
                        'Key Algorithm': cert.get('KeyAlgorithm', ''),
                    })
        except Exception as e:
            logger.debug(f"ACM {region}: {e}")
    return rows


def get_opensearch(credentials, regions):
    rows = []
    for region in regions:
        try:
            oss = _make_client(credentials, 'opensearch', region)
            domains = oss.list_domain_names().get('DomainNames', [])
            if not domains:
                continue
            details = oss.describe_domains(
                DomainNames=[d['DomainName'] for d in domains]
            ).get('DomainStatusList', [])
            for d in details:
                cluster_cfg = d.get('ClusterConfig', {})
                deployment = 'Multi-AZ' if cluster_cfg.get('ZoneAwarenessEnabled') else 'Single-AZ'
                engine_ver = d.get('EngineVersion', '')
                engine, version = (engine_ver.split('_', 1) + [''])[:2] if '_' in engine_ver else (engine_ver, '')
                rows.append({
                    'Region': region,
                    'Name': d.get('DomainName', ''),
                    'Engine': engine,
                    'Version': version,
                    'Deployment': deployment,
                    'Endpoint': d.get('Endpoint') or next(iter(d.get('Endpoints', {}).values()), ''),
                })
        except Exception as e:
            logger.debug(f"OpenSearch {region}: {e}")
    return rows


def get_api_gateways(credentials, regions):
    rows = []
    for region in regions:
        try:
            # REST APIs (v1)
            apigw = _make_client(credentials, 'apigateway', region)
            for page in apigw.get_paginator('get_rest_apis').paginate():
                for api in page['items']:
                    rows.append({
                        'Region': region,
                        'Name': api.get('name', ''),
                        'ID': api.get('id', ''),
                        'Protocol': 'REST',
                        'API Endpoint Type': ', '.join(
                            api.get('endpointConfiguration', {}).get('types', [])
                        ),
                    })
        except Exception as e:
            logger.debug(f"APIGW REST {region}: {e}")
        try:
            # HTTP / WebSocket APIs (v2)
            apigwv2 = _make_client(credentials, 'apigatewayv2', region)
            for page in apigwv2.get_paginator('get_apis').paginate():
                for api in page['Items']:
                    rows.append({
                        'Region': region,
                        'Name': api.get('Name', ''),
                        'ID': api.get('ApiId', ''),
                        'Protocol': api.get('ProtocolType', ''),
                        'API Endpoint Type': 'REGIONAL',
                    })
        except Exception as e:
            logger.debug(f"APIGW v2 {region}: {e}")
    return rows


def get_lambda_functions(credentials, regions):
    rows = []
    for region in regions:
        try:
            lmb = _make_client(credentials, 'lambda', region)
            for page in lmb.get_paginator('list_functions').paginate():
                for fn in page['Functions']:
                    rows.append({
                        'Region': region,
                        'Function Name': fn.get('FunctionName', ''),
                        'Runtime': fn.get('Runtime', 'N/A'),
                    })
        except Exception as e:
            logger.debug(f"Lambda {region}: {e}")
    return rows


def get_dynamodb_tables(credentials, regions):
    rows = []
    for region in regions:
        try:
            ddb = _make_client(credentials, 'dynamodb', region)
            for page in ddb.get_paginator('list_tables').paginate():
                for table_name in page['TableNames']:
                    rows.append({'Region': region, 'Name': table_name})
        except Exception as e:
            logger.debug(f"DynamoDB {region}: {e}")
    return rows


def get_sns_topics(credentials, regions):
    rows = []
    for region in regions:
        try:
            sns = _make_client(credentials, 'sns', region)
            for page in sns.get_paginator('list_topics').paginate():
                for topic in page['Topics']:
                    arn = topic['TopicArn']
                    topic_type = 'FIFO' if arn.endswith('.fifo') else 'Standard'
                    rows.append({
                        'Region': region,
                        'Name': arn.split(':')[-1],
                        'Type': topic_type,
                        'ARN': arn,
                    })
        except Exception as e:
            logger.debug(f"SNS {region}: {e}")
    return rows


# ---------------------------------------------------------------------------
# Account processing
# ---------------------------------------------------------------------------

def assume_role(session, account_id):
    sts = session.client('sts')
    resp = sts.assume_role(
        RoleArn=f"arn:aws:iam::{account_id}:role/{ROLE_NAME}",
        RoleSessionName=SESSION_NAME,
        DurationSeconds=3600,
    )
    return resp['Credentials']


def process_account(account, session, regions):
    account_id = account['Id']
    account_name = account.get('Name', account_id)
    logger.info(f"Processing account: {account_name} ({account_id})")

    try:
        caller_id = session.client('sts').get_caller_identity()['Account']
        credentials = session.get_credentials().get_frozen_credentials()
        credentials = {
            'AccessKeyId': credentials.access_key,
            'SecretAccessKey': credentials.secret_key,
            'SessionToken': credentials.token,
        } if account_id == caller_id else assume_role(session, account_id)
    except Exception as e:
        logger.error(f"Could not get credentials for {account_id}: {e}")
        return None

    result = {
        'account_id': account_id,
        'account_name': account_name,
        'vpcs': get_vpcs(credentials, regions),
        'ec2': get_ec2_instances(credentials, regions),
        'elb': get_load_balancers(credentials, regions),
        's3': get_s3_buckets(credentials),
        'elasticache': get_elasticache(credentials, regions),
        'rds': get_rds(credentials, regions),
        'acm': get_acm_certificates(credentials, regions),
        'opensearch': get_opensearch(credentials, regions),
        'apigw': get_api_gateways(credentials, regions),
        'lambda': get_lambda_functions(credentials, regions),
        'dynamodb': get_dynamodb_tables(credentials, regions),
        'sns': get_sns_topics(credentials, regions),
    }
    logger.info(f"Completed account: {account_name} ({account_id})")
    return result


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

SHEETS = {
    'VPCs':         ('vpcs',        ['Account ID', 'Account Name', 'Region', 'Name', 'VPC ID']),
    'EC2':          ('ec2',         ['Account ID', 'Account Name', 'Region', 'Name', 'Instance Type', 'Availability Zone', 'Platform Details']),
    'Load Balancers':('elb',        ['Account ID', 'Account Name', 'Region', 'Name', 'State', 'Type', 'VPC ID', 'Availability Zones']),
    'S3 Buckets':   ('s3',         ['Account ID', 'Account Name', 'Name', 'AWS Region']),
    'ElastiCache':  ('elasticache', ['Account ID', 'Account Name', 'Region', 'Cache Name', 'Status', 'Engine Version', 'Configuration']),
    'RDS':          ('rds',         ['Account ID', 'Account Name', 'Region', 'DB Identifier', 'Status', 'Role', 'Engine', 'Region & AZ', 'Size', 'Multi-AZ']),
    'ACM':          ('acm',         ['Account ID', 'Account Name', 'Region', 'Certificate ID', 'Domain Name', 'Type', 'Status', 'In Use', 'Renewal Eligibility', 'Key Algorithm']),
    'OpenSearch':   ('opensearch',  ['Account ID', 'Account Name', 'Region', 'Name', 'Engine', 'Version', 'Deployment', 'Endpoint']),
    'API Gateway':  ('apigw',       ['Account ID', 'Account Name', 'Region', 'Name', 'ID', 'Protocol', 'API Endpoint Type']),
    'Lambda':       ('lambda',      ['Account ID', 'Account Name', 'Region', 'Function Name', 'Runtime']),
    'DynamoDB':     ('dynamodb',    ['Account ID', 'Account Name', 'Region', 'Name']),
    'SNS':          ('sns',         ['Account ID', 'Account Name', 'Region', 'Name', 'Type', 'ARN']),
}

HEADER_FILL = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
HEADER_FONT = Font(color='FFFFFF', bold=True)


def _write_sheet(wb, sheet_name, headers, rows):
    ws = wb.create_sheet(title=sheet_name)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center')
    for row in rows:
        ws.append([row.get(h, '') for h in headers])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(
            len(str(cell.value or '')) for cell in col
        ) + 4
    return ws


def export_to_excel(all_results, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    aggregated = {key: [] for key in SHEETS}

    for result in all_results:
        if not result:
            continue
        acct_id = result['account_id']
        acct_name = result['account_name']
        for sheet_name, (data_key, _) in SHEETS.items():
            for row in result[data_key]:
                aggregated[sheet_name].append({'Account ID': acct_id, 'Account Name': acct_name, **row})

    for sheet_name, (_, headers) in SHEETS.items():
        _write_sheet(wb, sheet_name, headers, aggregated[sheet_name])

    wb.save(output_path)
    logger.info(f"Results saved to: {output_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    log_path = os.path.expanduser('~/Documents')
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    create_logger(__name__, log_path, f'resource_discovery_{timestamp}.log')

    profile = input("AWS SSO profile name: ").strip() or 'default'
    session = boto3.Session(profile_name=profile, region_name='us-east-1')
    validate_sso_token(session)

    logger.info("Fetching account list...")
    # account_list = get_org_account_list(session)
    account_list = target_accounts
    logger.info(f"Found {len(account_list)} accounts")

    logger.info("Fetching enabled regions...")
    regions = get_regions(session.client('ec2', region_name='us-east-1'))
    logger.info(f"Regions: {regions}")

    all_results = []
    with ThreadPoolExecutor(max_workers=5) as executor:
        futures = {
            executor.submit(process_account, acct, session, regions): acct
            for acct in account_list
        }
        for future in as_completed(futures):
            acct = futures[future]
            try:
                all_results.append(future.result())
            except Exception as e:
                logger.error(f"Account {acct['Id']} failed: {e}")

    output_file = os.path.join(log_path, f'aws_resource_inventory_{timestamp}.xlsx')
    export_to_excel(all_results, output_file)


if __name__ == '__main__':
    main()
