#!/usr/bin/env python3
"""
Convert Markdown documentation to professionally formatted Word document.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn

def setup_styles(doc):
    """Configure professional document styles."""
    # Title style
    title_style = doc.styles.add_style('Custom Title', WD_STYLE_TYPE.PARAGRAPH)
    title_font = title_style.font
    title_font.name = 'Calibri'
    title_font.size = Pt(24)
    title_font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_after = Pt(18)
    
    # Heading styles
    for level in range(1, 4):
        style_name = f'Custom Heading {level}'
        heading_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        heading_font = heading_style.font
        heading_font.name = 'Calibri'
        heading_font.size = Pt(16 - level * 2)
        heading_font.bold = True
        heading_style.paragraph_format.space_before = Pt(12)
        heading_style.paragraph_format.space_after = Pt(6)
    
    # Code style
    code_style = doc.styles.add_style('Code Block', WD_STYLE_TYPE.PARAGRAPH)
    code_font = code_style.font
    code_font.name = 'Consolas'
    code_font.size = Pt(9)
    code_style.paragraph_format.left_indent = Inches(0.5)
    #code_style.paragraph_format.space_before = Pt(6)
    #code_style.paragraph_format.space_after = Pt(6)

def add_table_borders(table):
    """Add borders to table."""
    tbl = table._tbl
    for row in tbl.tr_lst:
        for cell in row.tc_lst:
            tcPr = cell.tcPr
            tcBorders = OxmlElement('w:tcBorders')
            for border_name in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '4')
                border.set(qn('w:space'), '0')
                border.set(qn('w:color'), '000000')
                tcBorders.append(border)
            tcPr.append(tcBorders)

def convert_markdown_to_word(md_file_path, output_path):
    """Convert Markdown file to Word document."""
    doc = Document()
    setup_styles(doc)
    
    # Set document margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    with open(md_file_path, 'r', encoding='utf-8') as file:
        content = file.read()
    
    lines = content.split('\n')
    in_code_block = False
    code_content = []
    in_table = False
    table_rows = []
    
    for line in lines:
        # Handle code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                if code_content:
                    code_text = '\n'.join(code_content)
                    p = doc.add_paragraph(code_text, style='Code Block')
                    # Add light gray background
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:fill'), 'F5F5F5')
                    p._element.get_or_add_pPr().append(shd)
                code_content = []
                in_code_block = False
            else:
                in_code_block = True
            continue
        
        if in_code_block:
            code_content.append(line)
            continue
        
        # Handle tables
        if '|' in line and line.strip():
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append([cell.strip() for cell in line.split('|')[1:-1]])
            continue
        elif in_table:
            # End of table
            if len(table_rows) > 1:
                table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
                table.style = 'Table Grid'
                
                for i, row_data in enumerate(table_rows):
                    if i == 1 and all(cell.strip().startswith('-') for cell in row_data):
                        continue  # Skip separator row
                    
                    row_idx = i if i == 0 else i - 1
                    if row_idx >= len(table.rows):
                        continue
                        
                    for j, cell_data in enumerate(row_data):
                        if j < len(table.rows[row_idx].cells):
                            cell = table.rows[row_idx].cells[j]
                            cell.text = cell_data
                            if i == 0:  # Header row
                                for paragraph in cell.paragraphs:
                                    for run in paragraph.runs:
                                        run.bold = True
                
                add_table_borders(table)
            in_table = False
            table_rows = []
        
        # Handle headings
        if line.startswith('#'):
            level = len(line) - len(line.lstrip('#'))
            text = line.lstrip('# ').strip()
            
            if level == 1:
                doc.add_paragraph(text, style='Custom Title')
            elif level <= 3:
                doc.add_paragraph(text, style=f'Custom Heading {level}')
            else:
                p = doc.add_paragraph(text)
                p.runs[0].bold = True
        
        # Handle bold text
        elif '**' in line:
            p = doc.add_paragraph()
            parts = re.split(r'\*\*(.*?)\*\*', line)
            for i, part in enumerate(parts):
                if i % 2 == 0:
                    p.add_run(part)
                else:
                    p.add_run(part).bold = True
        
        # Handle inline code
        elif '`' in line:
            p = doc.add_paragraph()
            parts = re.split(r'`([^`]+)`', line)
            for i, part in enumerate(parts):
                if i % 2 == 0:
                    p.add_run(part)
                else:
                    run = p.add_run(part)
                    run.font.name = 'Consolas'
                    run.font.size = Pt(10)
        
        # Handle bullet points
        elif line.strip().startswith('- '):
            text = line.strip()[2:]
            p = doc.add_paragraph(text, style='List Bullet')
        
        # Regular paragraph
        elif line.strip():
            doc.add_paragraph(line.strip())
    
    # Save document
    doc.save(output_path)
    print(f"Word document saved to: {output_path}")

if __name__ == "__main__":
    # Input and output paths
    input_file = r"c:\Users\sgwalls\Documents\AWS_Projects\Scripts\Python\Git_repository\Walls-AWS\Identity_Center_Migration_Glossary.md"
    output_file = r"c:\Users\sgwalls\Documents\AWS_Projects\Scripts\Python\Git_repository\Walls-AWS\Identity_Center_Migration_Glossary.docx"
    
    # Convert the file
    convert_markdown_to_word(input_file, output_file)